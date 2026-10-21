#!/usr/bin/env python3
"""Wealth & Migration Summit 2026 — A4 letterhead (.docx).

Header : the Summit's horizontal blue lockup (left), a gold hairline under it.
Footer : a mist hairline, then three zones — title (left) · Global Citizen Solutions' official one-line mark
         in GCS blue (centre) · the Summit's website (right, a live link). Repeats on every page.
Body   : empty and ready to type in, or filled from a `letter` block in the spec.

    python3 letterhead.py build  [spec.json] [--out DIR] [--name N] [--title T] [--url U]
                                 [--fonts office|brand] [--no-rule] [--no-preview]
    python3 letterhead.py check  FILE.docx [--preview-dir DIR] [--no-render]

Spec (all keys optional; command-line flags win over the spec):

    {
      "name": "wm-letterhead",                         # file name, no extension
      "title": "Wealth and Migration Summit 2026",     # footer, left (shown as a tracked uppercase label)
      "url": "wealthandmigrationsummit.com",           # footer, right; linked to https://<url>
      "fonts": "office",                               # office = Cambria + Arial (installed everywhere)
                                                       # brand  = Caladea + Archivo, embedded in the file (Word, Regular only)
      "header_rule": true,                             # the gold hairline under the header logo
      "letter": {                                      # omit for a blank letterhead
        "date": "20 October 2026",
        "recipient": ["Name", "Company", "Address line"],
        "subject": "Subject line",
        "salutation": "Dear Name,",
        "body": ["Paragraph one.", "Paragraph two."],
        "closing": "Yours sincerely,",
        "signatory": ["Name", "Role"]
      }
    }

Only needs python-docx. `check` renders through LibreOffice (soffice) and poppler (pdfinfo, pdftotext,
pdfimages, pdftoppm) when they are on PATH, and says so when they are not.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor

SKILL = Path(__file__).resolve().parent.parent
LOGOS = SKILL / "assets" / "logos"
HEADER_LOGO = LOGOS / "wm-horizontal-blue.png"
FOOTER_LOGO = LOGOS / "GCS-Secondary-Blue.png"
FONTS = SKILL / "assets" / "fonts"

# ── Brand (site DESIGN.md / PRODUCT.md) ────────────────────────────────────────────────────────────
NAVY, MIST, MIST_DIM, GOLD = "121B31", "C8CCD7", "636B7C", "C2A358"   # gold only as a rule, never as text
DEFAULT_TITLE = "Wealth and Migration Summit 2026"
DEFAULT_URL = "wealthandmigrationsummit.com"
FONT_SETS = {
    # Cambria ships with every Office; Caladea (the brand serif) is its metric-compatible twin.
    # Arial is the nearest safe grotesk to Archivo. Default: the file looks the same on any machine.
    "office": {"head": "Cambria", "body": "Arial"},
    "brand": {"head": "Caladea", "body": "Archivo"},
}
EMBED = {"Caladea": "Caladea-Regular.ttf", "Archivo": "Archivo-400.ttf"}   # OFL; Regular only, bold is synthesised

# ── Page geometry, millimetres (A4 = 210 x 297) ────────────────────────────────────────────────────
PAGE_W, PAGE_H = 210, 297
MARGIN_X = 24                 # left / right (GCS house letterhead: 23.8); header and footer share the text edges
MARGIN_TOP = 40               # where the body starts
MARGIN_BOTTOM = 28            # room for the footer (it needs ~16 mm above FOOTER_DIST)
HEADER_DIST, FOOTER_DIST = 12, 10
HEADER_LOGO_W = 80            # horizontal lockup, as supplied (13.2 : 1)
FOOTER_COLS = (58, 54, 50)    # title | GCS mark | url  — must sum to PAGE_W - 2 * MARGIN_X
GCS_LOGO_W = 46               # keep >= 6 mm narrower than its column: LibreOffice clips a picture that nearly fills its cell
assert sum(FOOTER_COLS) == PAGE_W - 2 * MARGIN_X

ALT = {"header": "Wealth and Migration Summit", "footer": "Global Citizen Solutions"}


# ── XML helpers ────────────────────────────────────────────────────────────────────────────────────
def _el(tag: str, **attrs: str):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def set_fonts(rpr_owner, name: str) -> None:
    """Pin ascii/hAnsi/cs/eastAsia so no theme font can override the chosen one."""
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def style_run(run, font: str, size: float, color: str, bold: bool = False, caps: bool = False,
              tracking_em: float = 0.0) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    set_fonts(run._r, font)
    rpr = run._r.get_or_add_rPr()
    if caps:
        run.font.all_caps = True
    if tracking_em:   # character spacing, twentieths of a point; schema order puts it right before w:sz
        rpr.find(qn("w:sz")).addprevious(_el("w:spacing", w_val=round(tracking_em * size * 20)))


def tight(par, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    pf = par.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    par.alignment = align


def add_picture(par, path: Path, width_mm: float, alt: str):
    run = par.add_run()
    pic = run.add_picture(str(path), width=Mm(width_mm))
    inline = pic._inline
    for side in ("distT", "distB", "distL", "distR"):   # python-docx omits these; LibreOffice then insets the picture ~3 mm
        inline.set(side, "0")
    doc_pr = inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", alt)
    return pic


def add_link(par, url: str, text: str, font: str, size: float, color: str) -> None:
    href = url if re.match(r"^[a-z][a-z0-9+.-]*:", url) else f"https://{url}"
    rid = par.part.relate_to(href, RT.HYPERLINK, is_external=True)
    link = _el("w:hyperlink", r_id=rid, w_history="1")
    run = par.add_run(text)
    style_run(run, font, size, color)
    link.append(run._r)          # moves the run inside the hyperlink
    par._p.append(link)


# ── Header / footer ────────────────────────────────────────────────────────────────────────────────
def build_header(section, rule: bool) -> None:
    header = section.header
    header.is_linked_to_previous = False
    par = header.paragraphs[0]
    tight(par)
    add_picture(par, HEADER_LOGO, HEADER_LOGO_W, ALT["header"])
    if rule:
        pbdr = _el("w:pBdr")
        pbdr.append(_el("w:bottom", w_val="single", w_sz=4, w_space=9, w_color=GOLD))   # 0.5 pt gold hairline
        par._p.get_or_add_pPr().find(qn("w:spacing")).addprevious(pbdr)   # schema order: pBdr before spacing


def build_footer(section, title: str, url: str, fonts: dict) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    text_w = section.page_width - section.left_margin - section.right_margin
    tbl = footer.add_table(1, 3, text_w)
    tbl.autofit = False

    tbl_pr = tbl._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblInd", "w:tblBorders", "w:tblCellMar", "w:tblLook"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    tbl_pr.append(_el("w:tblW", w_w=Mm(sum(FOOTER_COLS)).twips, w_type="dxa"))
    tbl_pr.append(_el("w:tblInd", w_w=0, w_type="dxa"))
    borders = _el("w:tblBorders")
    borders.append(_el("w:top", w_val="single", w_sz=4, w_space=0, w_color=MIST))        # 0.5 pt mist hairline
    tbl_pr.append(borders)
    tbl_pr.append(_el("w:tblLayout", w_type="fixed"))
    mar = _el("w:tblCellMar")
    for side, w in (("top", 140), ("left", 0), ("bottom", 0), ("right", 0)):
        mar.append(_el(f"w:{side}", w_w=w, w_type="dxa"))
    tbl_pr.append(mar)

    for col, w in zip(tbl.columns, FOOTER_COLS):
        col.width = Mm(w)
    left, mid, right = tbl.rows[0].cells
    for cell, w in zip((left, mid, right), FOOTER_COLS):
        cell.width = Mm(w)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = left.paragraphs[0]
    tight(p)
    style_run(p.add_run(title), fonts["body"], 6.5, NAVY, caps=True, tracking_em=0.16)   # a label: uppercase, 0.16em

    p = mid.paragraphs[0]
    tight(p, WD_ALIGN_PARAGRAPH.CENTER)
    add_picture(p, FOOTER_LOGO, GCS_LOGO_W, ALT["footer"])

    p = right.paragraphs[0]
    tight(p, WD_ALIGN_PARAGRAPH.RIGHT)
    add_link(p, url, re.sub(r"^[a-z][a-z0-9+.-]*://", "", url).rstrip("/"), fonts["body"], 7.5, MIST_DIM)   # shown without the scheme

    # A footer must end on a paragraph: reuse the one python-docx created, collapsed to nothing.
    tail = footer.paragraphs[0]
    tbl._tbl.addnext(tail._p)
    tight(tail)
    tail.paragraph_format.line_spacing = Pt(1)
    tail.paragraph_format.line_spacing_rule = 4   # WD_LINE_SPACING.EXACTLY
    tail.add_run().font.size = Pt(1)


# ── Font embedding (brand fonts) ───────────────────────────────────────────────────────────────────
def embed_fonts(path: Path, families=EMBED) -> None:
    """Embed the brand TTFs the way Word does: obfuscated (.odttf), listed in fontTable.xml, flagged in settings."""
    from lxml import etree
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    table = etree.fromstring(parts["word/fontTable.xml"])
    rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    for i, (family, filename) in enumerate(families.items(), 1):
        guid = uuid.uuid4()
        key = bytes.fromhex(guid.hex)[::-1]                       # the GUID, reversed, is the XOR key
        data = (FONTS / filename).read_bytes()
        parts[f"word/fonts/font{i}.odttf"] = bytes(b ^ key[n % 16] for n, b in enumerate(data[:32])) + data[32:]
        rels.append(f'<Relationship Id="rId{i}" Type="{r}/font" Target="fonts/font{i}.odttf"/>')
        for old in table.findall(f"{{{w}}}font[@{{{w}}}name='{family}']"):
            table.remove(old)
        font = etree.SubElement(table, f"{{{w}}}font", {f"{{{w}}}name": family})
        etree.SubElement(font, f"{{{w}}}charset", {f"{{{w}}}val": "00"})
        etree.SubElement(font, f"{{{w}}}family", {f"{{{w}}}val": "swiss" if family == "Archivo" else "roman"})
        etree.SubElement(font, f"{{{w}}}pitch", {f"{{{w}}}val": "variable"})
        etree.SubElement(font, f"{{{w}}}embedRegular", {f"{{{r}}}id": f"rId{i}", f"{{{w}}}fontKey": "{" + str(guid).upper() + "}"})
    rels.append("</Relationships>")
    parts["word/fontTable.xml"] = etree.tostring(table, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts["word/_rels/fontTable.xml.rels"] = "".join(rels).encode()

    ct = parts["[Content_Types].xml"].decode()
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ', 1)
    parts["[Content_Types].xml"] = ct.encode()

    settings = etree.fromstring(parts["word/settings.xml"])
    if settings.find(f"{{{w}}}embedTrueTypeFonts") is None:
        flag = etree.Element(f"{{{w}}}embedTrueTypeFonts")
        zoom = settings.find(f"{{{w}}}zoom")
        (zoom.addnext(flag) if zoom is not None else settings.insert(0, flag))
    parts["word/settings.xml"] = etree.tostring(settings, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", parts.pop("[Content_Types].xml"))      # must be first
        for name, blob in parts.items():
            z.writestr(name, blob)
    tmp.replace(path)


# ── Document ───────────────────────────────────────────────────────────────────────────────────────
def setup_document(fonts: dict, title: str):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(PAGE_W), Mm(PAGE_H)
    s.left_margin = s.right_margin = Mm(MARGIN_X)
    s.top_margin, s.bottom_margin = Mm(MARGIN_TOP), Mm(MARGIN_BOTTOM)
    s.header_distance, s.footer_distance = Mm(HEADER_DIST), Mm(FOOTER_DIST)

    styles = doc.styles
    normal = styles["Normal"]
    set_fonts(normal.element, fonts["body"])
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    def para_style(name: str, font: str, size: float, before: float = 0, after: float = 0, bold: bool = False):
        st = styles.add_style(name, 1)
        st.base_style = normal
        set_fonts(st.element, font)
        st.font.size, st.font.bold = Pt(size), bold
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
        return st

    para_style("Letter Address", fonts["body"], 10)
    para_style("Letter Subject", fonts["head"], 13, before=10, after=8)

    # Word opens a python-docx file in "compatibility mode" (14); 15 is the current layout engine.
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))        # the template's bare zoom fails the schema (percent is required)
    if zoom is not None:
        zoom.attrib.clear()
        zoom.set(qn("w:percent"), "100")
    for cs in settings.iter(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            cs.set(qn("w:val"), "15")
    for lang in list(doc.styles.element.iter(qn("w:lang"))) + list(settings.iter(qn("w:themeFontLang"))):
        lang.set(qn("w:val"), "en-GB")

    cp = doc.core_properties
    cp.title, cp.subject = f"Letterhead — {title}", "A4 letterhead"
    cp.author = cp.last_modified_by = "Global Citizen Solutions"
    cp.keywords, cp.comments, cp.language = "letterhead; Wealth and Migration Summit 2026", "", "en-GB"
    return doc


def add_letter(doc, letter: dict) -> None:
    """Pour a letter into the body. Text goes in verbatim — nothing is added, reworded or invented."""

    def add(text: str, style: str | None = None, after: float | None = None, before: float | None = None,
            keep_next: bool = False, bold: bool = False):
        p = doc.add_paragraph(style=style)
        run = p.add_run(text)
        run.font.bold = True if bold else None
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.keep_with_next = keep_next or None
        return p

    if letter.get("date"):
        add(letter["date"], after=14)
    recipient = letter.get("recipient") or []
    for i, line in enumerate(recipient):
        add(line, "Letter Address", after=10 if i == len(recipient) - 1 else 0)
    if letter.get("subject"):
        add(letter["subject"], "Letter Subject", keep_next=True)
    if letter.get("salutation"):
        add(letter["salutation"], after=8, keep_next=True)
    for para in letter.get("body") or []:
        add(para)
    if letter.get("closing"):
        add(letter["closing"], before=8, keep_next=True)
    signatory = letter.get("signatory") or []
    for i, line in enumerate(signatory):
        add(line, "Letter Address", before=34 if i == 0 else None, keep_next=i < len(signatory) - 1, bold=i == 0)


def build(spec: dict, out_dir: Path) -> Path:
    title = spec.get("title") or DEFAULT_TITLE
    url = spec.get("url") or DEFAULT_URL
    fonts = FONT_SETS[spec.get("fonts", "office")]
    doc = setup_document(fonts, title)
    section = doc.sections[0]
    build_header(section, spec.get("header_rule", True))
    build_footer(section, title, url, fonts)
    if spec.get("letter"):
        add_letter(doc, spec["letter"])
    else:
        doc.add_paragraph()      # an empty line to type into
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{spec.get('name') or 'wm-letterhead'}.docx"
    doc.save(path)
    if spec.get("fonts", "office") == "brand":
        embed_fonts(path)
    return path


# ── Check ──────────────────────────────────────────────────────────────────────────────────────────
def _text(el) -> str:
    """Visible text of a header/footer/body element; paragraphs separated by a space."""
    paras = list(el.iter(qn("w:p")))
    return re.sub(r"\s+", " ", " ".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in paras)).strip()


def _sha(b: bytes) -> str:
    return hashlib.sha1(b).hexdigest()


def _images(part) -> list[tuple[str, str]]:
    """(sha1 of the embedded file, alt text) for every picture in a header/footer part."""
    out = []
    for blip in part._element.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        doc_pr = next(blip.iterancestors(qn("wp:inline"))).find(qn("wp:docPr"))
        out.append((_sha(part.part.related_parts[rid].blob), doc_pr.get("descr") or ""))
    return out


def check(path: Path, render: bool = True, preview_dir: Path | None = None) -> int:
    errors: list[str] = []
    warns: list[str] = []
    notes: list[str] = []
    doc = Document(str(path))

    # A4, portrait
    for i, s in enumerate(doc.sections, 1):
        w, h = s.page_width.twips, s.page_height.twips
        if abs(w - 11906) > 3 or abs(h - 16838) > 3:
            errors.append(f"section {i}: page is {w}x{h} twips, not A4 portrait (11906x16838)")
        if s.orientation != 0:
            errors.append(f"section {i}: not portrait")
        text_w = s.page_width - s.left_margin - s.right_margin
        notes.append(f"section {i}: A4, margins L/R {s.left_margin.mm:.0f} mm, top {s.top_margin.mm:.0f}, "
                     f"bottom {s.bottom_margin.mm:.0f}; text width {Emu(text_w).mm:.0f} mm")
        if s.bottom_margin.mm < s.footer_distance.mm + 16:
            warns.append(f"section {i}: bottom margin {s.bottom_margin.mm:.0f} mm leaves under 16 mm for the footer "
                         "— Word will push the body up")
        if s.header.is_linked_to_previous or s.footer.is_linked_to_previous:
            if i > 1:
                notes.append(f"section {i}: inherits header/footer from the previous section")

        want_h = _sha(HEADER_LOGO.read_bytes())
        want_f = _sha(FOOTER_LOGO.read_bytes())
        if i == 1 or not s.header.is_linked_to_previous:
            hi = _images(s.header)
            if len(hi) != 1:
                errors.append(f"section {i} header: expected the one Summit logo, found {len(hi)} picture(s)")
            elif hi[0][0] != want_h:
                errors.append(f"section {i} header: the picture is not the supplied wm-horizontal-blue logo")
            elif not hi[0][1]:
                errors.append(f"section {i} header: logo has no alt text")
        if i == 1 or not s.footer.is_linked_to_previous:
            fi = _images(s.footer)
            if len(fi) != 1:
                errors.append(f"section {i} footer: expected the one Global Citizen Solutions mark, found {len(fi)}")
            elif fi[0][0] != want_f:
                errors.append(f"section {i} footer: the picture is not the supplied GCS-Secondary-Blue mark")
            elif not fi[0][1]:
                errors.append(f"section {i} footer: GCS mark has no alt text")
            ftext = _text(s.footer._element)
            if not ftext.strip():
                errors.append(f"section {i} footer: no text")
            for tc in s.footer._element.iter(qn("w:tc")):
                cell_w = int(tc.find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w"))) / 1440 * 25.4
                for ext in tc.iter(qn("wp:extent")):
                    pic_w = int(ext.get("cx")) / 36000
                    if pic_w > cell_w - 6:
                        errors.append(f"section {i} footer: a {pic_w:.0f} mm picture sits in a {cell_w:.0f} mm cell — "
                                      "LibreOffice clips it; widen the column or shrink the picture")
            first_cell = next(s.footer._element.iter(qn("w:tc")), None)
            if first_cell is not None and len(_text(first_cell)) > 34:
                warns.append(f"section {i} footer: the title is {len(_text(first_cell))} characters — over ~34 it wraps to "
                             f"two lines in its {FOOTER_COLS[0]} mm column")
            links = list(s.footer._element.iter(qn("w:hyperlink")))
            if not links:
                warns.append(f"section {i} footer: the website is not a link")
            if "@" in ftext:
                notes.append("footer carries an e-mail address — confirm it is the one meant to be public")

    # fonts: only the brand sets, explicit everywhere
    used = set()
    ours = [st for st in doc.styles.element.iter(qn("w:style")) if st.get(qn("w:styleId")) in ("Normal", "LetterAddress", "LetterSubject")]
    ours += list(doc.styles.element.iter(qn("w:rPrDefault")))
    for part in [doc.element] + ours + [x for s in doc.sections for x in (s.header._element, s.footer._element)]:
        for rf in part.iter(qn("w:rFonts")):
            used.update(v for k, v in rf.attrib.items() if k.endswith(("ascii", "hAnsi")) and "Theme" not in k)
    allowed = {f for fs in FONT_SETS.values() for f in fs.values()} | {"Times New Roman", "Symbol", "Calibri", "Cambria Math"}
    odd = sorted(used - allowed)
    if odd:
        warns.append(f"fonts outside the brand sets: {', '.join(odd)}")
    notes.append(f"fonts: {', '.join(sorted(used & allowed - {'Times New Roman', 'Symbol', 'Cambria Math'}))}")
    if used & set(EMBED):
        with zipfile.ZipFile(path) as z:
            n_embedded = len([n for n in z.namelist() if n.startswith("word/fonts/") and n.endswith(".odttf")])
        if n_embedded < len(used & set(EMBED)):
            warns.append("brand fonts are used but not embedded — they will only show where installed")
        else:
            notes.append(f"{n_embedded} brand font(s) embedded (Regular; bold is synthesised)")

    # render
    soffice, have = shutil.which("soffice"), {t: shutil.which(t) for t in ("pdfinfo", "pdftotext", "pdfimages", "pdftoppm")}
    if render and not soffice:
        warns.append("LibreOffice (soffice) not found — skipped the render. Open the file in Word and look at it.")
    elif render and all(have.values()):
        _render_checks(path, doc, preview_dir or path.parent, errors, warns, notes, soffice)
    elif render:
        warns.append("poppler tools missing (" + ", ".join(t for t, v in have.items() if not v) + ") — skipped the render")

    print(f"\n{path.name}")
    for n in notes:
        print(f"  ·  {n}")
    for w in warns:
        print(f"  WARN   {w}")
    for e in errors:
        print(f"  ERROR  {e}")
    print("  ✓ checks passed" if not errors else f"  ✗ {len(errors)} error(s)")
    return 1 if errors else 0


def _render_checks(path, doc, preview_dir, errors, warns, notes, soffice) -> None:
    footer = doc.sections[0].footer._element
    link = next(footer.iter(qn("w:hyperlink")), None)
    url = "".join(t.text or "" for t in link.iter(qn("w:t"))) if link is not None else ""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        proc = subprocess.run(
            [soffice, f"-env:UserInstallation={(tmp / 'lo').as_uri()}", "--headless", "--convert-to", "pdf",
             "--outdir", str(tmp), str(path)], capture_output=True, text=True, timeout=180)
        pdf = tmp / f"{path.stem}.pdf"
        if not pdf.exists():
            warns.append(f"render failed: {proc.stderr.strip()[-200:] or proc.stdout.strip()[-200:]}")
            return
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
        size = re.search(r"Page size:\s+([\d.]+) x ([\d.]+)", info)
        pw, ph = float(size.group(1)), float(size.group(2))
        if abs(pw - 595.3) > 2 or abs(ph - 841.9) > 2:
            errors.append(f"render: page is {pw:.0f}x{ph:.0f} pt, not A4 (595x842)")
        notes.append(f"render: {pages} page(s), A4")
        listing = subprocess.run(["pdfimages", "-list", str(pdf)], capture_output=True, text=True).stdout.splitlines()[2:]
        per_page: dict[int, int] = {}
        for row in listing:
            cols = row.split()
            if cols and cols[0].isdigit():
                per_page[int(cols[0])] = per_page.get(int(cols[0]), 0) + 1
        for pg in range(1, pages + 1):
            txt = subprocess.run(["pdftotext", "-f", str(pg), "-l", str(pg), "-layout", str(pdf), "-"],
                                 capture_output=True, text=True).stdout
            flat = re.sub(r"\s+", " ", txt).lower()
            if url and url.lower() not in flat:
                errors.append(f"render p.{pg}: website '{url}' missing from the footer")
            if per_page.get(pg, 0) < 2:
                errors.append(f"render p.{pg}: expected the header logo and the GCS mark, found {per_page.get(pg, 0)} picture(s)")
        preview_dir.mkdir(parents=True, exist_ok=True)
        stem = preview_dir / f"{path.stem}_preview"
        subprocess.run(["pdftoppm", "-r", "110", "-png", str(pdf), str(stem)], check=False)
        for p in sorted(preview_dir.glob(f"{path.stem}_preview*.png")):
            notes.append(f"preview: {p}")


# ── CLI ────────────────────────────────────────────────────────────────────────────────────────────
def default_out(name: str) -> Path:
    mnt = Path("/mnt/user-data/outputs")
    return mnt if mnt.is_dir() else Path.cwd() / "output" / name


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="build the .docx, then check and preview it")
    b.add_argument("spec", nargs="?", help="spec.json (optional)")
    b.add_argument("--out", help="output directory")
    b.add_argument("--name"), b.add_argument("--title"), b.add_argument("--url")
    b.add_argument("--fonts", choices=sorted(FONT_SETS))
    b.add_argument("--no-rule", action="store_true", help="drop the gold hairline under the header")
    b.add_argument("--no-preview", action="store_true", help="skip the LibreOffice render and PNG preview")

    c = sub.add_parser("check", help="check an existing letterhead .docx")
    c.add_argument("file")
    c.add_argument("--preview-dir")
    c.add_argument("--no-render", action="store_true")

    args = ap.parse_args()
    if args.cmd == "check":
        return check(Path(args.file), not args.no_render, Path(args.preview_dir) if args.preview_dir else None)

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8")) if args.spec else {}
    for key in ("name", "title", "url", "fonts"):
        if getattr(args, key):
            spec[key] = getattr(args, key)
    if args.no_rule:
        spec["header_rule"] = False
    out = Path(args.out) if args.out else default_out(spec.get("name") or "wm-letterhead")
    path = build(spec, out)
    print(f"wrote {path} ({path.stat().st_size // 1024} KB)")
    return check(path, not args.no_preview)


if __name__ == "__main__":
    sys.exit(main())
