#!/usr/bin/env python3
"""Lint and render a Wealth & Migration Summit deck.

    python3 scripts/check_deck.py lint   deck.pptx [--as-of YYYY-MM-DD] [--confirmed url,email]
    python3 scripts/check_deck.py render deck.pptx outdir        # PNG per slide + contact.png
    python3 scripts/check_deck.py all    deck.pptx outdir

`lint` reads the .pptx itself (so it also catches edits made in PowerPoint) and checks:

  FACTS    retired facts (old day themes, "150–200", a dinner, "register your interest"), urgency
           wording, prices not on the site or not in force on --as-of, speaker / workshop counts,
           the 200 cap, session times the site never published, unconfirmed URL / e-mail / CNN mark,
           claims the Summit cannot make yet (testimonials, past-edition figures, sponsor tiers)
  LAYOUT   text taller than its box (measured with the real Caladea/Archivo metrics), anything off
           the slide, type below 12pt, an unfilled placeholder, a slide with no title
  CONTRAST WCAG AA for text on its ground (gold on a light ground fails — gold is for dark grounds)
  FONTS    only theme fonts, or Cambria/Arial/Caladea/Archivo

Exit status 1 when there is any ERROR. WARNs are for the hand-off. `render` needs LibreOffice (soffice)
and pdftoppm; it loads Caladea/Archivo from the skill into a throw-away profile (Cambria renders as Caladea, its metric twin; Arial as Liberation Sans), so the PNGs wrap like PowerPoint — but `lint` is still the authority on fit.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
import facts  # noqa: E402
import wm_deck as wm  # noqa: E402

EMU_IN = 914400
ALLOWED_FONTS = {"Cambria", "Arial", "Caladea", "Archivo"}

RETIRED = [
    (r"golden visa pathway", "retired Day One theme — it is now “The New Geography of Wealth & Mobility”"),
    (r"structuring,\s*context\s*(&|and)\s*experience", "retired Day Two theme — it is now “Global Wealth & the International Family”"),
    (r"150\s*[–-]\s*200|\b150\b\s+(curated\s+)?(delegates|attendees|seats)", "the site says “Limited to 200”, never 150"),
    (r"networking dinner|private dinner|gala dinner|\bdinner\b", "there is no dinner — networking lunches both days, a cocktail reception on Day Two evening"),
    (r"register (your )?interest|registration of interest", "the call to action is “Secure your place” (a booking), not “register your interest”"),
]
URGENCY = r"last chance|hurry|don.t miss|only \d+\s+(seats?|places?|spots?|tickets?)|selling (fast|out)|sold out|book now|act now|limited time|few (seats|places) left"
ABSENT = r"testimonial|last year|previous edition|past edition|inaugural edition|sponsor tier|platinum sponsor|gold sponsor|official partner|returning delegates"


class Report:
    def __init__(self):
        self.items: list[tuple[str, int, str]] = []

    def add(self, level, slide, msg):
        self.items.append((level, slide, msg))

    @property
    def errors(self):
        return [i for i in self.items if i[0] == "ERROR"]

    def print(self):
        for level, n, msg in sorted(self.items, key=lambda i: (i[0] != "ERROR", i[1])):
            where = f"slide {n}" if n else "deck"
            print(f"  {level:<5} {where}: {msg}")


def _frames(shape):
    """Yield (shape, text_frame) for every text-bearing shape, descending into groups and tables."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            yield from _frames(s)
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        yield shape, shape.text_frame
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield shape, cell.text_frame


def slide_text(slide) -> str:
    out = [tf.text for sh in slide.shapes for _, tf in _frames(sh)]
    if slide.has_notes_slide:
        out.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def _bg_hex(slide) -> str | None:
    try:
        if slide.background.fill.type is not None:
            return str(slide.background.fill.fore_color.rgb)
    except (AttributeError, TypeError):
        pass
    return wm.PAL["navy"]  # the master maps bg1 to navy


def _overlaps(a, b) -> bool:
    ax2, ay2, bx2, by2 = a.left + a.width, a.top + a.height, b.left + b.width, b.top + b.height
    ix = max(0, min(ax2, bx2) - max(a.left, b.left))
    iy = max(0, min(ay2, by2) - max(a.top, b.top))
    return ix * iy > 0.3 * a.width * a.height


def _ground(shape, slide, bg: str) -> str:
    """The colour a text box really sits on: the last solid, opaque shape under it that contains it, else the slide."""
    ground = bg
    for other in slide.shapes:
        if other is shape or other.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        if other._element is shape._element:
            break
        try:
            if other.fill.type == 1 and other.width >= shape.width * 0.95 and other.height >= shape.height * 0.95 \
                    and other.left <= shape.left + 9144 and other.top <= shape.top + 9144 \
                    and other.left + other.width >= shape.left + shape.width - 9144 \
                    and other.top + other.height >= shape.top + shape.height - 9144 \
                    and other._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).find(qn("a:alpha")) is None:
                ground = str(other.fill.fore_color.rgb)
        except (AttributeError, TypeError, ValueError):
            continue
    return ground


def lint_layout(prs, rep: Report):
    sw, sh_ = prs.slide_width, prs.slide_height
    for n, slide in enumerate(prs.slides, 1):
        if slide.shapes.title is None or not slide.shapes.title.text_frame.text.strip():
            rep.add("WARN", n, "no title — the outline view and screen readers show this slide as untitled")
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width > Emu(int(sw * 0.4))]
        bg = _bg_hex(slide)
        for shape in slide.shapes:
            if shape.left is not None and (shape.left < -Emu(9144) or shape.top < -Emu(9144)
                                           or shape.left + shape.width > sw + Emu(9144)
                                           or shape.top + shape.height > sh_ + Emu(9144)):
                rep.add("ERROR", n, f"“{shape.name}” runs off the slide")
            if shape.is_placeholder and getattr(shape, "has_text_frame", False) and shape.has_text_frame \
                    and not shape.text_frame.text.strip():
                rep.add("ERROR", n, f"empty placeholder “{shape.name}” — it shows prompt text in PowerPoint")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and not shape._element.nvPicPr.cNvPr.get("descr") \
                    and not shape.name.lower().startswith("decor"):
                rep.add("WARN", n, f"picture “{shape.name}” has no alt text")
            for shp, tf in _frames(shape):
                _check_frame(shp, tf, n, _ground(shp, slide, bg), pictures, rep)


def _check_frame(shape, tf, n, bg, pictures, rep: Report):
    if not tf.text.strip():
        return
    w_in = (shape.width - (tf.margin_left or 0) - (tf.margin_right or 0)) / EMU_IN
    h_in = shape.height / EMU_IN
    need = 0.0
    over_photo = any(_overlaps(shape, p) for p in pictures)
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        pt = max((r.font.size.pt for r in runs if r.font.size), default=None)
        for r in runs:
            if r.font.size and r.font.size.pt < wm.MIN_PT:
                rep.add("ERROR", n, f"“{r.text[:30]}” is {r.font.size.pt:g}pt — nothing below {wm.MIN_PT}pt")
            name = r.font.name
            if name and not name.startswith("+") and name not in ALLOWED_FONTS:
                rep.add("WARN", n, f"font “{name}” on “{r.text[:24]}” — use the theme fonts")
            _check_contrast(r, bg, n, rep, over_photo)
        if pt is None:
            continue
        kind = "head" if any((r.font.name or "") == "+mj-lt" for r in runs) else "body"
        rp = runs[0]._r.rPr
        caps = rp is not None and rp.get("cap") == "all"
        track = (int(rp.get("spc", 0)) / 100.0) if rp is not None else 0.0
        text = "".join(r.text for r in runs)
        lines = wm.count_lines(text, kind, pt, w_in, track, caps)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        need += lines * pt * wm.LINE_K * ls / 72.0
        if p.space_after:
            need += p.space_after.pt / 72.0
    if tf.word_wrap is not False and need > h_in + 0.06:
        rep.add("ERROR", n, f"“{shape.name}”: text needs {need:.2f}in, the box is {h_in:.2f}in — it overflows")


def _check_contrast(run, bg, n, rep: Report, over_photo: bool):
    try:
        fg = str(run.font.color.rgb)
    except (AttributeError, TypeError):
        return
    if over_photo or bg is None:
        return
    big = bool(run.font.size and run.font.size.pt >= 24)
    ratio = wm.contrast(fg, bg)
    if ratio < (3.0 if big else 4.5):
        rep.add("ERROR", n, f"“{run.text[:30]}” is {fg} on {bg}: contrast {ratio:.1f}:1 fails AA"
                + (" — gold is for dark grounds, never text on a light one" if fg.upper() == wm.PAL["gold"] else ""))


def lint_facts(prs, rep: Report, as_of: str, confirmed: set):
    d = date.fromisoformat(as_of)
    if d > facts.EVENT_END:
        rep.add("ERROR", 0, f"as-of {as_of} is after the Summit (6 Nov 2026) — every fact here is forward-looking")
    for msg in facts.freshness(as_of):
        rep.add("WARN", 0, msg)
    known_times = facts.known_times()
    published = facts.published_amounts()
    inforce = facts.rate_on(as_of)
    roster = len(facts.speakers())
    workshops = sum(len(x["workshops"]) for x in facts.programme())
    whole = []
    for n, slide in enumerate(prs.slides, 1):
        text = slide_text(slide)
        whole.append(text)
        for pat, why in RETIRED:
            if re.search(pat, text, re.I):
                rep.add("ERROR", n, why)
        if re.search(URGENCY, text, re.I):
            rep.add("ERROR", n, "urgency wording — state the deadline as a date, never the pressure")
        if re.search(ABSENT, text, re.I):
            rep.add("ERROR", n, "claims the Summit cannot make yet (testimonials, past editions, sponsor tiers, official partners) — nothing on the site supports them")
        is_rates = any(s.name.startswith("Price ") for s in slide.shapes)
        prices = [int(p.replace(",", "")) for p in re.findall(r"€\s?(\d[\d,]*)", text)]
        for p in prices:
            if p not in published:
                rep.add("ERROR", n, f"€{p} is not a published rate (€{', €'.join(str(x) for x in sorted(published))})")
        if is_rates and inforce is None:
            rep.add("ERROR", n, f"a rates slide on {as_of}, when no published tier is in force (tiers run 1 Sep – 4 Nov 2026) — leave the slide out")
        if prices and not is_rates:
            if inforce is None:
                rep.add("ERROR", n, f"a price on {as_of}, when no published tier is in force — drop it")
            elif any(p != inforce["eur"] for p in prices if p in published):
                rep.add("ERROR", n, f"price shown is not the {inforce['name']} (€{inforce['eur']}), the tier in force on {as_of}")
        for m in re.finditer(r"(\d+)\s+speakers", text, re.I):
            if int(m.group(1)) != roster:
                rep.add("ERROR", n, f"“{m.group(0)}” — the site lists {roster} speakers")
        for m in re.finditer(r"(\d+)\s+workshops", text, re.I):
            if int(m.group(1)) != workshops:
                rep.add("ERROR", n, f"“{m.group(0)}” — the programme has {workshops} workshops (5 a day)")
        for m in re.finditer(r"limited to (\d+)", text, re.I):
            if int(m.group(1)) != 200:
                rep.add("ERROR", n, f"“{m.group(0)}” — the cap is 200")
        for m in re.finditer(r"\b(?:[01]\d|2[0-3]):[0-5]\d\b", text):
            if m.group(0) not in known_times:
                rep.add("WARN", n, f"{m.group(0)} is not a time the site publishes")
    joined = "\n".join(whole)
    if "wealthandmigrationsummit.com" in joined.lower() and "url" not in confirmed:
        rep.add("WARN", 0, "the Summit URL is printed but not confirmed live (the site sat behind a staging password) — confirm, then pass --confirmed url")
    if re.search(r"events@globalcitizensolutions\.com", joined, re.I) and "email" not in confirmed:
        rep.add("WARN", 0, "events@globalcitizensolutions.com is printed but not confirmed as the partner route — confirm, then pass --confirmed email")
    venues = {v for v in ("Marques", "Marquês") if re.search(rf"EPIC SANA {v}", joined)}
    if len(venues) > 1:
        rep.add("WARN", 0, "the venue is spelled both “Marques” and “Marquês” — the site does both; pick one (default: EPIC SANA Marques Hotel)")


def lint_bookends(prs, rep: Report):
    """The cover and the back cover are dark blue and carry the GCS and CNN Portugal marks (the user's decision)."""
    slides = list(prs.slides)
    if len(slides) < 2:
        return
    for n, label in ((1, "cover"), (len(slides), "back cover")):
        s = slides[n - 1]
        bg = _bg_hex(s)
        if bg and wm.luminance(bg) > 0.2:
            rep.add("WARN", n, f"the {label} is on a light ground — the cover and the back cover are dark blue")
        names = {sh.name.lower() for sh in s.shapes}
        missing = [m for m in ("gcs", "cnn") if f"logo {m}" not in names]
        if missing:
            rep.add("WARN", n, f"the {label} has no {' / '.join(m.upper() for m in missing)} mark — the site hero and the social cover pair them")


def lint_theme(prs, rep: Report):
    # a hand-edited theme drifting from the brand is worth one line
    root = prs.slide_master.part.part_related_by("http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    blob = root.blob.decode("utf-8", "ignore")
    for tag, expect in (("accent1", wm.PAL["gold"]), ("dk1", wm.PAL["navy"])):
        m = re.search(rf"<a:{tag}><a:srgbClr val=\"([0-9A-Fa-f]{{6}})\"", blob)
        if m and m.group(1).upper() != expect.upper():
            rep.add("WARN", 0, f"theme {tag} is {m.group(1)}, the brand value is {expect}")


def run(path, as_of=None, confirmed=None, spec=None) -> int:
    prs = Presentation(str(path))
    rep = Report()
    as_of = as_of or date.today().isoformat()
    lint_facts(prs, rep, as_of, set(confirmed or []))
    lint_layout(prs, rep)
    lint_bookends(prs, rep)
    lint_theme(prs, rep)
    n_err = len(rep.errors)
    print(f"check {Path(path).name}: {len(prs.slides)} slides, as of {as_of} — "
          f"{n_err} error(s), {sum(1 for i in rep.items if i[0] == 'WARN')} warning(s)")
    rep.print()
    return 1 if n_err else 0


def render(path, outdir) -> list[Path]:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not Path(soffice).exists() and not shutil.which("soffice"):
        print("soffice (LibreOffice) not found — install it to render, or open the .pptx in PowerPoint")
        return []
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        # A throw-away LibreOffice profile that carries the brand fonts: LibreOffice maps Cambria to Caladea
        # on its own, so with Caladea present the preview wraps exactly like PowerPoint's Cambria would.
        fonts = Path(td) / "profile" / "user" / "fonts"
        fonts.mkdir(parents=True)
        for ttf in (wm.ASSETS / "fonts").glob("*.ttf"):
            shutil.copy(ttf, fonts / ttf.name)
        subprocess.run([soffice, f"-env:UserInstallation={(Path(td) / 'profile').as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", td, str(path)], check=True, capture_output=True, timeout=300)
        pdf = Path(td) / (Path(path).stem + ".pdf")
        for old in out.glob("slide-*.png"):
            old.unlink()
        subprocess.run(["pdftoppm", "-png", "-r", "72", str(pdf), str(out / "slide")], check=True, capture_output=True, timeout=240)
        shutil.copy(pdf, out / (Path(path).stem + ".pdf"))
    pngs = sorted(out.glob("slide-*.png"))
    if pngs:
        from PIL import Image
        thumbs = [Image.open(p) for p in pngs]
        tw, th = 480, round(480 * thumbs[0].height / thumbs[0].width)
        cols = 3
        rows = -(-len(thumbs) // cols)
        sheet = Image.new("RGB", (cols * (tw + 10) + 10, rows * (th + 10) + 10), (60, 60, 60))
        for i, im in enumerate(thumbs):
            sheet.paste(im.resize((tw, th)), (10 + (i % cols) * (tw + 10), 10 + (i // cols) * (th + 10)))
        sheet.save(out / "contact.png")
    print(f"rendered {len(pngs)} slide(s) → {out}  (Cambria shown as Caladea, its metric twin)")
    return pngs


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("cmd", choices=["lint", "render", "all"])
    ap.add_argument("deck")
    ap.add_argument("outdir", nargs="?", default=None)
    ap.add_argument("--as-of")
    ap.add_argument("--confirmed", default="", help="comma list of facts the user confirmed: url,email")
    a = ap.parse_args(argv)
    code = 0
    if a.cmd in ("lint", "all"):
        code = run(a.deck, a.as_of, {c for c in a.confirmed.split(",") if c})
    if a.cmd in ("render", "all"):
        render(a.deck, a.outdir or (Path(a.deck).with_suffix("")).as_posix() + "-preview")
    return code


if __name__ == "__main__":
    sys.exit(main())
