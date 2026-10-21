#!/usr/bin/env python3
"""Wealth & Migration Summit 2026 — PowerPoint deck builder.

Turns a JSON spec into an editable .pptx in the Summit's "Lisbon Nocturne" identity:
navy ground, amber light, Caladea-style serif headings over a grotesk body, hairlines
instead of boxes, square corners. Facts (programme, speakers, rate tiers, the 200 cap)
are PULLED from references/data/*.json — a spec names what to show, it never retypes
a time, price or name.

    python3 scripts/wm_deck.py spec.json out.pptx [--fonts office|brand] [--as-of YYYY-MM-DD]

Slide types (see references/deck-design.md for what each is for):
    cover · statement · split · photo · ledger · entries · days · programme ·
    speakers · speaker · rates · section · closing

Needs: python-pptx, Pillow, lxml. Nothing else, no network.

As a library:
    sys.path.insert(0, "<SKILL_DIR>/scripts"); import wm_deck
    deck = wm_deck.Deck(fonts="office", as_of="2026-10-20")
    deck.add(wm_deck.SlideSpec(...))   # or deck.build(spec_dict)
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SKILL_DIR = Path(__file__).resolve().parent.parent
ASSETS = SKILL_DIR / "assets"
sys.path.insert(0, str(Path(__file__).resolve().parent))
import facts  # noqa: E402

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

# ----------------------------------------------------------------------------
# Palette — read from DESIGN.md so a newer export re-colours every future deck
# ----------------------------------------------------------------------------

_DEFAULT_COLORS = {
    "navy": "121B31", "navy-deep": "0A101F", "navy-wash": "12192D", "dawn-blue": "232F52",
    "gold": "C2A358", "gold-bright": "DCC08B", "dawn-low": "1B2643", "mist": "C8CCD7", "white": "FFFFFF",
}


def load_palette() -> dict:
    colors = dict(_DEFAULT_COLORS)
    try:
        front = (SKILL_DIR / "references" / "DESIGN.md").read_text(encoding="utf-8").split("---")[1]
    except (OSError, IndexError):
        return colors
    in_colors = False
    for line in front.splitlines():
        if line.startswith("colors:"):
            in_colors = True
            continue
        if in_colors:
            m = re.match(r'\s+([a-z-]+):\s*"#([0-9A-Fa-f]{6})"', line)
            if m:
                colors[m.group(1)] = m.group(2).upper()
            elif line.strip() and not line.startswith(" "):
                break
    return colors


PAL = load_palette()


def mix(a: str, b: str, t: float) -> str:
    """Blend hex `a` toward hex `b` by t (0..1) — a hairline is mist at low alpha, pre-mixed onto the ground."""
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


def luminance(hexc: str) -> float:
    def ch(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hexc[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast(a: str, b: str) -> float:
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


# tone = the ground a slide sits on. The Hour rule: night deepens toward the programme, lifts toward the close.
TONES = {
    "paper": "FFFFFF",  # a white slide: the user allows white slides interleaved with the night ones
    "navy": PAL["navy"],
    "deep": PAL["navy-deep"],
    "dawn": PAL["dawn-blue"],
}

# ----------------------------------------------------------------------------
# Type — the closed ramp, at slide scale (pt). Names match DESIGN.md's roles.
# ----------------------------------------------------------------------------

FONT_SETS = {
    # Cambria ships with every Office; Caladea (the brand serif) is its metric-compatible twin, so the
    # layout is identical. Arial is the nearest safe grotesk to Archivo. This is the default: it
    # opens correctly on any machine.
    "office": {"head": "Cambria", "body": "Arial"},
    # The real brand fonts (bundled in assets/fonts) — only correct where they are installed.
    "brand": {"head": "Caladea", "body": "Archivo"},
}
MEASURE_FONTS = {"head": ASSETS / "fonts" / "Caladea-Regular.ttf", "body": ASSETS / "fonts" / "Archivo-400.ttf"}

ROLES = {
    #            font   pt   line-spacing  caps   tracking(em)
    "display":  ("head", 48, 0.95, False, 0.0),
    "headline": ("head", 36, 1.00, False, 0.0),
    "title":    ("head", 28, 1.05, False, 0.0),
    "subtitle": ("head", 22, 1.08, False, 0.0),
    "session":  ("head", 18, 1.15, False, 0.0),
    "lead":     ("body", 20, 1.30, False, 0.0),
    "body":     ("body", 16, 1.35, False, 0.0),
    "datum":    ("body", 14, 1.30, False, 0.0),
    "label":    ("body", 12, 1.00, True, 0.16),
}
MIN_PT = 12  # the audience skews older: nothing below it

# ----------------------------------------------------------------------------
# Geometry (inches) — 16:9, one shared gutter
# ----------------------------------------------------------------------------

W, H = 13.333, 7.5
MX = 0.8                 # side gutter
CW = W - 2 * MX          # content width
TOP = 0.75               # first line of content
FOOT_Y = 6.95            # footer baseline box
BOTTOM = 6.6             # lowest point content may reach
LINE_K = 1.2             # font-size → line-height factor used for measuring

_font_cache: dict = {}


def _font(kind: str, pt: float) -> ImageFont.FreeTypeFont:
    key = (kind, pt)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(str(MEASURE_FONTS[kind]), int(round(pt * 10)))
    return _font_cache[key]


def text_width_pt(text: str, kind: str, pt: float, track_pt: float = 0.0) -> float:
    return _font(kind, pt).getlength(text) / 10.0 + track_pt * len(text)


def count_lines(text: str, kind: str, pt: float, width_in: float, track_pt: float = 0.0, caps: bool = False) -> int:
    # 4% safety: PowerPoint sets text without kerning, PIL measures with it, and renderers differ by a few px
    maxw = width_in * 72.0 * 0.96
    total = 0
    for para in str(text).split("\n"):
        para = para.upper() if caps else para
        cur, n = "", 1
        for word in para.split(" "):
            trial = f"{cur} {word}" if cur else word
            if not cur or text_width_pt(trial, kind, pt, track_pt) <= maxw:
                cur = trial
            else:
                n += 1
                cur = word
        total += n
    return total


def text_height_in(text: str, role: str, width_in: float, ls: float | None = None) -> float:
    kind, pt, rls, caps, track = ROLES[role]
    n = count_lines(text, kind, pt, width_in, pt * track, caps)
    return n * pt * LINE_K * (ls or rls) / 72.0


def initials(name: str) -> str:
    parts = [p for p in re.sub(r"\(.*?\)|,.*$|\b(H\.E\.|Dr\.)", "", name).split() if p[0].isalpha()]
    return (parts[0][0] + parts[-1][0]).upper() if len(parts) > 1 else parts[0][:2].upper()


# ----------------------------------------------------------------------------
# Low-level shape helpers
# ----------------------------------------------------------------------------

def rgb(hexc: str) -> RGBColor:
    return RGBColor.from_string(hexc.upper())


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _alpha(fill_parent, alpha: float):
    """alpha 0..1 on the srgbClr inside a solidFill."""
    clr = fill_parent.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


class Slide:
    """Thin wrapper so section builders read like layout code."""

    def __init__(self, deck: "Deck", pptx_slide, number: int):
        self.deck, self.s, self.n = deck, pptx_slide, number

    # --- shapes ----------------------------------------------------------
    def rect(self, x, y, w, h, fill, alpha=None, name="Block"):
        shp = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        _strip_style(shp)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _alpha(shp._element.spPr, alpha)
        shp.line.fill.background()
        shp.name = name
        return shp

    def oval(self, x, y, d, fill, name="Station dot"):
        shp = self.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        _strip_style(shp)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        shp.line.fill.background()
        shp.name = name
        return shp

    def scrim(self, x, y, w, h, color, a0, a1, angle=0, name="Scrim"):
        """Rectangle with a colour gradient from alpha a0 to a1 (0..1) at `angle` (0 = left→right, 90 = top→bottom)."""
        shp = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        _strip_style(shp)
        shp.line.fill.background()
        spPr = shp._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
            for el in spPr.findall(qn(tag)):
                spPr.remove(el)
        grad = parse_xml(
            f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>'
            f'<a:gs pos="0"><a:srgbClr val="{color}"><a:alpha val="{int(a0 * 100000)}"/></a:srgbClr></a:gs>'
            f'<a:gs pos="100000"><a:srgbClr val="{color}"><a:alpha val="{int(a1 * 100000)}"/></a:srgbClr></a:gs>'
            f'</a:gsLst><a:lin ang="{int(angle * 60000)}" scaled="0"/></a:gradFill>')
        spPr.find(qn("a:prstGeom")).addnext(grad)
        shp.name = name
        return shp

    def hline(self, x1, y, x2, color=None, pt=0.75, name="Rule"):
        return self._line(x1, y, x2, y, color or self.rule(), pt, name)

    def vline(self, x, y1, y2, color=None, pt=0.75, name="Rule"):
        return self._line(x, y1, x, y2, color or self.rule(), pt, name)

    def _line(self, x1, y1, x2, y2, color, pt, name):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        _strip_style(c)
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(pt)
        c.name = name
        return c

    # --- text ------------------------------------------------------------
    def text(self, x, y, w, h, paras, role="body", color=None, align="l", anchor="t",
             ls=None, after=0, name="Text", link=None):
        box = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        self.deck.fill_tf(box.text_frame, paras, role, color or self.ink, align, anchor, ls, after, link)
        self.deck.check_fit(self, name, paras, role, w, h, ls, after)
        return box

    def title(self, text, role="headline", x=MX, y=TOP, w=CW, h=None, color=None, align="l"):
        """Fill the layout's real title placeholder, so the outline view and screen readers get a title."""
        shp = self.s.shapes.title
        need = text_height_in(text, role, w)
        h = h or need + 0.05
        shp.left, shp.top, shp.width, shp.height = Inches(x), Inches(y), Inches(w), Inches(h)
        self.deck.fill_tf(shp.text_frame, text, role, color or self.ink, align, "t", None, 0, None)
        self.deck.check_fit(self, "Title", text, role, w, h, None, 0)
        tree = self.s.shapes._spTree  # the title placeholder is created first; lift it above photos and scrims
        tree.remove(shp._element)
        tree.append(shp._element)
        if role in ("display", "headline") and count_lines(text, ROLES[role][0], ROLES[role][1], w) > 3:
            self.deck.warn(self.n, f"title runs to more than 3 lines at {ROLES[role][1]}pt — shorten it or widen the box")
        return shp, h

    def picture(self, path, x, y, w, h, alt="", focus=(0.5, 0.5), name="Picture"):
        img, (iw, ih) = self.deck.prep_image(path)
        pic = self.s.shapes.add_picture(str(img), Inches(x), Inches(y), Inches(w), Inches(h))
        target, src = w / h, iw / ih
        if src > target:
            extra = 1 - target / src
            pic.crop_left, pic.crop_right = extra * focus[0], extra * (1 - focus[0])
        elif src < target:
            extra = 1 - src / target
            pic.crop_top, pic.crop_bottom = extra * focus[1], extra * (1 - focus[1])
        pic._element.nvPicPr.cNvPr.set("descr", alt)
        pic.name = name
        return pic

    def logo(self, key, x, y, w=None, h=None, alt=""):
        path = self.deck.logo(key)
        iw, ih = Image.open(path).size
        if w is None:
            w = h * iw / ih
        else:
            h = w * ih / iw
        pic = self.s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        pic._element.nvPicPr.cNvPr.set("descr", alt)
        pic.name = f"Logo {key}"
        return pic, w, h

    def notes(self, text):
        if text:
            self.s.notes_slide.notes_text_frame.text = text

    @property
    def ground(self) -> str:
        return self.deck.tone_hex[self.n]

    # Colours by ground. On the night grounds: white ink, mist for the quiet text, gold for the one lit thing.
    # On paper (a white slide): navy ink, navy at 66% for the quiet text, and gold stays a MARK only — the lamp, a
    # dot, the programme line — never text (DESIGN.md: gold never touches paper; it fails AA as text).
    @property
    def paper(self) -> bool:
        return self.ground == TONES["paper"]

    @property
    def ink(self) -> str:
        return PAL["navy"] if self.paper else PAL["white"]

    @property
    def soft(self) -> str:
        return mix(TONES["paper"], PAL["navy"], 0.66) if self.paper else PAL["mist"]

    @property
    def accent(self) -> str:
        return PAL["navy"] if self.paper else PAL["gold"]

    def rule(self, strong: bool = False) -> str:
        if self.paper:
            return mix(TONES["paper"], PAL["navy"], 0.32 if strong else 0.16)
        return mix(self.ground, PAL["mist"], 0.32 if strong else 0.20)


# ----------------------------------------------------------------------------
# The deck
# ----------------------------------------------------------------------------

class Deck:
    def __init__(self, fonts: str = "office", as_of: str | None = None, lang: str = "en-GB",
                 title: str = "Wealth & Migration Summit 2026", author: str = "Global Citizen Solutions",
                 confirmed: list | None = None):
        if fonts not in FONT_SETS:
            raise ValueError(f"fonts must be one of {sorted(FONT_SETS)}")
        self.fonts, self.fontset, self.lang = fonts, FONT_SETS[fonts], lang
        self.as_of = as_of or date.today().isoformat()
        self.confirmed = set(confirmed or [])
        self.warnings: list[tuple[int, str]] = []
        self.tone_hex: dict[int, str] = {}
        self._tmp = tempfile.TemporaryDirectory(prefix="wm_deck_")
        self._img: dict = {}
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)
        self.prs.core_properties.title = title
        self.prs.core_properties.author = author
        self.prs.core_properties.subject = "Wealth & Migration Summit 2026 · Lisbon · 5–6 November 2026"
        self._theme()
        self._master_and_layouts()
        self.cover_layout = next(l for l in self.prs.slide_layouts if l.name == "WM Cover")
        self.content_layout = next(l for l in self.prs.slide_layouts if l.name == "WM Content")
        self.paper_layout = next(l for l in self.prs.slide_layouts if l.name == "WM Content Paper")

    # --- theme, master, layouts -----------------------------------------
    def _theme(self):
        master = self.prs.slide_master
        part = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        root.set("name", "Wealth & Migration Summit 2026")
        cs = root.find(f".//{{{A_NS}}}clrScheme")
        cs.set("name", "Lisbon Nocturne")
        scheme = {
            "dk1": PAL["navy"], "lt1": PAL["white"], "dk2": PAL["navy-deep"], "lt2": PAL["mist"],
            "accent1": PAL["gold"], "accent2": PAL["gold-bright"], "accent3": PAL["dawn-blue"],
            "accent4": PAL["mist"], "accent5": PAL["dawn-low"], "accent6": PAL["navy-wash"],
            "hlink": PAL["gold-bright"], "folHlink": PAL["mist"],
        }
        for tag, hexv in scheme.items():
            el = cs.find(f"{{{A_NS}}}{tag}")
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{{{A_NS}}}srgbClr").set("val", hexv)
        fs = root.find(f".//{{{A_NS}}}fontScheme")
        fs.set("name", "Lisbon Nocturne")
        fs.find(f"{{{A_NS}}}majorFont/{{{A_NS}}}latin").set("typeface", self.fontset["head"])
        fs.find(f"{{{A_NS}}}minorFont/{{{A_NS}}}latin").set("typeface", self.fontset["body"])
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _master_and_layouts(self):
        prs, master = self.prs, self.prs.slide_master
        # Dark mapping: the ground is dk1 (navy) and default text is lt1 (white), so a text box a colleague
        # adds in PowerPoint is readable instead of navy-on-navy.
        cm = master._element.find(qn("p:clrMap"))
        cm.set("bg1", "dk1"), cm.set("tx1", "lt1"), cm.set("bg2", "dk2"), cm.set("tx2", "lt2")
        # one title style: left, 36pt, theme heading font
        t = master._element.find(".//" + qn("p:titleStyle") + "/" + qn("a:lvl1pPr"))
        t.set("algn", "l")
        rpr = t.find(qn("a:defRPr"))
        rpr.set("sz", "3600")
        rpr.set("b", "0")
        for lay in list(prs.slide_layouts):
            if lay.name not in ("Title Slide", "Title Only", "Title and Content"):
                prs.slide_layouts.remove(lay)
        for sh in list(master.placeholders):
            if sh.placeholder_format.type in (13, 15, 16):  # date, footer, slide number
                sh._element.getparent().remove(sh._element)
        for lay in prs.slide_layouts:
            for sh in list(lay.placeholders):
                if sh.placeholder_format.type in (13, 15, 16):
                    sh._element.getparent().remove(sh._element)
            lay._element.cSld.set("name", {"Title Slide": "WM Cover", "Title Only": "WM Content",
                                           "Title and Content": "WM Content Paper"}[lay.name])
        # reposition master title/body for 16:9
        for sh in master.placeholders:
            if sh.placeholder_format.type == 1:
                sh.left, sh.top, sh.width, sh.height = Inches(MX), Inches(TOP), Inches(CW), Inches(1.2)
            elif sh.placeholder_format.type == 2:
                sh.left, sh.top, sh.width, sh.height = Inches(MX), Inches(2.0), Inches(CW), Inches(4.4)
        for lay in prs.slide_layouts:
            for sh in lay.placeholders:
                ph = sh.placeholder_format.type
                if ph in (1, 3):  # title / centred title
                    sh.left, sh.top, sh.width, sh.height = Inches(MX), Inches(TOP), Inches(CW), Inches(1.2)
                elif ph == 4:  # subtitle
                    sh.left, sh.top, sh.width, sh.height = Inches(MX), Inches(3.9), Inches(8.6), Inches(1.2)
        content = next(l for l in prs.slide_layouts if l.name == "WM Content")
        self._layout_footer(content, PAL["mist"])
        paper = next(l for l in prs.slide_layouts if l.name == "WM Content Paper")
        for sh in list(paper.placeholders):  # a title and nothing else, like WM Content
            if sh.placeholder_format.type not in (1, 3):
                sh._element.getparent().remove(sh._element)
        # A white layout: the colour map is flipped (bg1 = lt1, tx1 = dk1), so what a colleague adds in PowerPoint is
        # navy on white. Gold stays out of the text colours on this ground.
        paper._element.cSld.insert(0, parse_xml(
            f'<p:bg xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:bgPr><a:solidFill><a:schemeClr val="bg1"/></a:solidFill>'
            f'<a:effectLst/></p:bgPr></p:bg>'))
        self._paper_clrmap(paper._element)
        self._layout_footer(paper, mix(TONES["paper"], PAL["navy"], 0.66))

    @staticmethod
    def _paper_clrmap(root):
        """Replace a layout's or slide's colour-map override with the light mapping."""
        old = root.find(qn("p:clrMapOvr"))
        new = parse_xml(
            f'<p:clrMapOvr xmlns:p="{P_NS}" xmlns:a="{A_NS}"><a:overrideClrMapping bg1="lt1" tx1="dk1" bg2="lt2" '
            f'tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" '
            f'accent6="accent6" hlink="hlink" folHlink="folHlink"/></p:clrMapOvr>')
        if old is not None:
            old.addprevious(new)
            root.remove(old)
        else:
            root.append(new)

    def _layout_footer(self, layout, color):
        """Footer label + slide number live on the layout, never on a slide, so they restyle in one place."""
        tree = layout.shapes._spTree
        y, h = Inches(FOOT_Y), Inches(0.3)
        label_spc = int(ROLES["label"][1] * ROLES["label"][4] * 100)

        def sp(idn, name, x, w, algn, body):
            return parse_xml(
                f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="{idn}" name="{name}"/>'
                f'<p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
                f'<p:spPr><a:xfrm><a:off x="{int(Inches(x))}" y="{int(y)}"/><a:ext cx="{int(Inches(w))}" cy="{int(h)}"/></a:xfrm>'
                f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
                f'<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"><a:noAutofit/></a:bodyPr>'
                f'<a:lstStyle/><a:p><a:pPr algn="{algn}"/>{body}</a:p></p:txBody></p:sp>')

        rpr = (f'<a:rPr lang="en-GB" sz="1200" cap="all" spc="{label_spc}"><a:solidFill><a:srgbClr val="{color}"/>'
               f'</a:solidFill><a:latin typeface="+mn-lt"/></a:rPr>')
        tree.append(sp(901, "Footer label", MX, 6.0, "l", f'<a:r>{rpr}<a:t>Wealth &amp; Migration Summit 2026</a:t></a:r>'))
        tree.append(sp(902, "Slide number", W - MX - 1.5, 1.5, "r",
                       f'<a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">{rpr}<a:t>‹#›</a:t></a:fld>'))

    # --- assets ----------------------------------------------------------
    def prep_image(self, path):
        path = Path(path)
        if not path.is_absolute():
            path = ASSETS / path
        if path in self._img:
            return self._img[path]
        im = Image.open(path)
        size = im.size
        if path.suffix.lower() in (".jpg", ".jpeg", ".png"):
            out = (path, size)
        else:  # webp → jpeg, capped at 2000px wide so the deck stays light
            im = im.convert("RGB")
            if im.width > 2000:
                im = im.resize((2000, round(im.height * 2000 / im.width)), Image.LANCZOS)
            dst = Path(self._tmp.name) / (path.stem + ".jpg")
            im.save(dst, "JPEG", quality=86)
            out = (dst, im.size)
        self._img[path] = out
        return out

    LOGOS = {
        "wm": "w&m-default white.png", "wm-horizontal": "w&m-horizontal White.png", "wm-symbol": "w&m-symbol.png",
        "wm-vertical": "w&m-vertical white.png", "gcs": "GCS-Primary-White.png", "gcs-symbol": "GCS-Symbol-White.png",
        "cnn": "Logo CNN PORTUGAL Float BRANCO Final.png",
        # white grounds: GCS primary blue + CNN Portugal red (`Float Final`)
        "gcs-blue": "GCS-Primary-Blue.png", "cnn-red": "Logo CNN PORTUGAL Float Final.png",
    }

    def logo(self, key: str) -> Path:
        return ASSETS / "logos" / self.LOGOS[key]

    # --- text plumbing ---------------------------------------------------
    def warn(self, n: int, msg: str):
        self.warnings.append((n, msg))

    def fill_tf(self, tf, paras, role, color, align, anchor, ls, after, link):
        kind, pt, rls, caps, track = ROLES[role]
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, str):
            paras = [paras]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = ls or rls
            if after and i < len(paras) - 1:
                p.space_after = Pt(after)
            runs = para if isinstance(para, list) else [(para, {})]
            for text, o in runs:
                r = p.add_run()
                r.text = text
                f = r.font
                rk = ROLES[o.get("role", role)]
                f.size = Pt(rk[1])
                f.name = "+mj-lt" if rk[0] == "head" else "+mn-lt"
                f.bold = False
                f.color.rgb = rgb(o.get("color", color))
                rp = r._r.get_or_add_rPr()
                rp.set("lang", self.lang)
                if rk[3]:
                    rp.set("cap", "all")
                    rp.set("spc", str(int(rk[1] * rk[4] * 100)))
                if link:
                    r.hyperlink.address = link

    def check_fit(self, sl: Slide, name, paras, role, w, h, ls, after):
        if isinstance(paras, str):
            paras = [paras]
        need = 0.0
        for i, para in enumerate(paras):
            if isinstance(para, list):
                txt, prole = "".join(t for t, _ in para), para[0][1].get("role", role)
            else:
                txt, prole = para, role
            need += text_height_in(txt, prole, w, ls)
            if after and i < len(paras) - 1:
                need += after / 72.0
        if need > h + 0.04:
            self.warn(sl.n, f"{name}: text needs {need:.2f}in but the box is {h:.2f}in — cut words or split the slide")

    # --- spec resolution -------------------------------------------------
    @staticmethod
    def resolve(value):
        """'@site.lede' → the site's own line, verbatim. Anything else passes through."""
        if isinstance(value, str) and value.startswith("@"):
            src, _, key = value[1:].partition(".")
            data = {"site": facts.site, "rates": facts.rates}[src]()
            cur = data
            for part in key.split("."):
                cur = cur[int(part)] if isinstance(cur, list) else cur[part]
            return cur
        return value

    # --- slide factory ---------------------------------------------------
    def new_slide(self, tone: str, cover: bool = False) -> Slide:
        if tone not in TONES:
            raise ValueError(f"tone must be one of {sorted(TONES)}, got {tone!r}")
        if cover and tone == "paper":
            raise ValueError("the cover and the back cover are always on a dark ground")
        layout = self.cover_layout if cover else self.paper_layout if tone == "paper" else self.content_layout
        s = self.prs.slides.add_slide(layout)
        if tone == "paper":
            self._paper_clrmap(s._element)
        n = len(self.prs.slides)
        self.tone_hex[n] = TONES[tone]
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(TONES[tone])
        return Slide(self, s, n)

    def build(self, spec: dict):
        meta = spec.get("meta", {})
        alternate, k = bool(meta.get("alternate")), 0
        for slide_spec in spec["slides"]:
            kind = slide_spec["type"]
            if alternate and kind not in ("cover", "closing") and "tone" not in slide_spec:
                # night, paper, night, paper… — the programme keeps its deepest tone on the night slides
                slide_spec = dict(slide_spec, tone="paper" if k % 2 else ("deep" if kind == "programme" else "navy"))
                k += 1
            fn = getattr(self, f"_slide_{kind}", None)
            if fn is None:
                raise ValueError(f"unknown slide type {kind!r}; choose from {SLIDE_TYPES}")
            fn(slide_spec)
        self._report_programme_flags()
        return self

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return Path(path)

    # ======================================================================
    # Slide types
    # ======================================================================

    def _finish(self, sl: Slide, sp: dict):
        sl.notes(sp.get("notes"))
        # an unfilled subtitle placeholder shows "Click to add subtitle" in PowerPoint
        for ph in list(sl.s.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)
        return sl

    # --- cover -----------------------------------------------------------
    def _slide_cover(self, sp):
        sl = self.new_slide(sp.get("tone", "navy"), cover=True)
        image = sp.get("image")
        if image:
            path = {"hero": "img/hero-1920.webp", "venue": "img/epic-sana-1440.webp",
                    "night": "img/nocturne-texture-1440.webp"}.get(image, image)
            sl.picture(path, 0, 0, W, H, alt=sp.get("image_alt", "Lisbon at blue hour, with the Tagus"), name="Cover photograph")
            # The scrim is layered: one crossing from the text side, one rising from the ground.
            sl.scrim(0, 0, W, H, sl.ground, 0.94, 0.40, 0, "Scrim from the text side")
            sl.scrim(0, H * 0.55, W, H * 0.45, sl.ground, 0.0, 0.88, 90, "Scrim from the ground")
        sl.logo("wm-horizontal", MX, 0.62, w=3.5, alt="Wealth & Migration Summit")
        y = 1.7
        service = self.resolve(sp.get("service_line", "@site.service_line"))
        if service:
            sl.rect(MX, y + 0.075, 0.1, 0.1, PAL["gold"], name="Lamp")
            sl.text(MX + 0.28, y, 7.0, 0.25, service, "label", sl.accent, name="Service line")
            y += 0.55
        title = self.resolve(sp["title"])
        _, th = sl.title(title, "display", MX, y, 9.4, color=sl.ink)
        y += th + 0.3
        lede = self.resolve(sp.get("lede"))
        if lede:
            lh = text_height_in(lede, "lead", 8.4)
            sl.text(MX, y, 8.4, lh + 0.05, lede, "lead", sl.soft, name="Lede")
            y += lh
        if y > 6.3:
            self.warn(sl.n, f"cover stack ends at {y:.2f}in, into the logo zone — shorten the title or lede")
        self._marks(sl, sp.get("marks", ["gcs", "cnn"]))
        return self._finish(sl, sp)

    def _marks(self, sl: Slide, marks, y=6.5, h=0.55, gap=0.45):
        """Global Citizen Solutions + CNN Portugal, side by side, left-aligned at the foot — the pair the site
        hero shows and the social cover and closing carry. White on the dark grounds; on a white ground the GCS primary
        blue and the red CNN Portugal. The GCS mark stands alone: nothing written above, beside or over it (its name is
        in the mark)."""
        gcs, cnn = ("gcs-blue", "cnn-red") if sl.paper else ("gcs", "cnn")
        if sl.paper:
            y = min(y, 6.1)   # a white slide has the footer label at the foot (the covers have none): sit above it
        x = MX
        if "gcs" in marks:
            _, w, _ = sl.logo(gcs, x, y, h=h, alt="Global Citizen Solutions")
            x += w + gap
        if "cnn" in marks:
            sl.logo(cnn, x, y, h=h, alt="CNN Portugal")

    # --- statement -------------------------------------------------------
    def _slide_statement(self, sp):
        sl = self.new_slide(sp.get("tone", "navy"))
        has_image = bool(sp.get("image"))
        w = 7.2 if has_image else 9.6
        y = 1.15
        _, h = sl.title(self.resolve(sp["headline"]), "headline", MX, y, w)
        y += h + 0.4
        if sp.get("lead"):
            lead = self.resolve(sp["lead"])
            lh = text_height_in(lead, "lead", w)
            sl.text(MX, y, w, lh + 0.05, lead, "lead", sl.ink, name="Lead")
            y += lh + 0.3
        body = [self.resolve(b) for b in sp.get("body", [])]
        if body:
            bh = sum(text_height_in(b, "body", w) for b in body) + 0.14 * (len(body) - 1)
            sl.text(MX, y, w, bh + 0.05, body, "body", sl.soft, after=10, name="Body")
            y += bh
        if y > BOTTOM:
            self.warn(sl.n, f"content ends at {y:.2f}in, past the footer line — split into two slides")
        if has_image:
            sl.picture(sp["image"], 8.75, 1.15, W - MX - 8.75, 5.2, alt=sp.get("image_alt", ""), name="Plate")
        return self._finish(sl, sp)

    # --- split (photo bleeds off the right edge and dissolves into the ground) ---
    def _slide_split(self, sp):
        sl = self.new_slide(sp.get("tone", "navy"))
        px = 7.4
        sl.picture(sp["image"], px, 0, W - px, H, alt=sp.get("image_alt", ""), focus=tuple(sp.get("focus", (0.5, 0.5))), name="Plate")
        sl.scrim(px, 0, 2.2, H, sl.ground, 1.0, 0.0, 0, "Dissolve into the ground")
        y = 1.15
        w = 6.2
        _, h = sl.title(self.resolve(sp["headline"]), "headline", MX, y, w)
        y += h + 0.4
        body = [self.resolve(b) for b in sp.get("body", [])]
        if body:
            bh = sum(text_height_in(b, "body", w) for b in body) + 0.14 * (len(body) - 1)
            sl.text(MX, y, w, bh + 0.05, body, "body", sl.soft, after=10, name="Body")
            y += bh
        if sp.get("caption"):
            cap = sp["caption"]
            cw = min(text_width_pt(cap.upper(), "body", 12, 1.92) / 72 + 0.45, 4.5)
            sl.rect(px + 0.9, 6.15, cw, 0.4, PAL["navy-deep"], alpha=0.72, name="Caption chip")
            sl.text(px + 1.1, 6.15, cw - 0.4, 0.4, cap, "label", sl.soft, anchor="m", name="Caption")
        if y > BOTTOM:
            self.warn(sl.n, f"content ends at {y:.2f}in, past the footer line")
        return self._finish(sl, sp)

    # --- photo (full bleed, headline over the lower scrim) ---------------
    def _slide_photo(self, sp):
        sl = self.new_slide(sp.get("tone", "deep"))
        sl.picture(sp["image"], 0, 0, W, H, alt=sp.get("image_alt", ""), focus=tuple(sp.get("focus", (0.5, 0.5))), name="Photograph")
        sl.scrim(0, H * 0.35, W, H * 0.65, sl.ground, 0.0, 0.92, 90, "Scrim from the ground")
        sl.scrim(0, 0, W * 0.6, H, sl.ground, 0.55, 0.0, 0, "Scrim from the text side")
        text = self.resolve(sp["headline"])
        h = text_height_in(text, "headline", 8.6)
        sl.title(text, "headline", MX, 6.35 - h, 8.6)
        if sp.get("caption"):
            sl.text(MX, 6.45, 8, 0.25, sp["caption"], "label", sl.soft, name="Caption")
        return self._finish(sl, sp)

    # --- ledger (a few figures, hairline above each) ---------------------
    AUTO_FIGURES = {
        "speakers": (lambda: str(len(facts.speakers())), "speakers"),
        "workshops": (lambda: str(sum(len(d["workshops"]) for d in facts.programme())), "workshops"),
        "days": (lambda: str(len(facts.programme())), "days"),
        "capacity": (lambda: re.search(r"\d+", facts.rates()["limit"]).group(0), "attendees at most"),
    }

    def _slide_ledger(self, sp):
        sl = self.new_slide(sp.get("tone", "navy"))
        _, h = sl.title(self.resolve(sp["headline"]), "headline", MX, 1.0, 9.6)
        items = sp["items"]
        if not 2 <= len(items) <= 4:
            raise ValueError("ledger takes 2–4 items")
        gap = 0.45
        cw = (CW - gap * (len(items) - 1)) / len(items)
        y0 = 3.55
        for i, it in enumerate(items):
            if "auto" in it:
                fig, default_label = self.AUTO_FIGURES[it["auto"]]
                figure, label = fig(), it.get("label", default_label)
            else:
                figure, label = it["figure"], it["label"]
            x = MX + i * (cw + gap)
            sl.hline(x, y0, x + cw, color=sl.rule(True), name="Rule above figure")
            sl.text(x, y0 + 0.25, cw, 0.9, figure, "display", sl.accent, name=f"Figure {i + 1}")
            lh = text_height_in(label, "lead", cw)
            sl.text(x, y0 + 1.3, cw, lh + 0.05, label, "lead", sl.ink, name=f"Figure {i + 1} label")
        if sp.get("note"):
            sl.text(MX, 6.2, CW, 0.4, sp["note"], "datum", sl.soft, name="Note")
        return self._finish(sl, sp)

    # --- entries (ruled columns, no boxes) --------------------------------
    def _slide_entries(self, sp):
        sl = self.new_slide(sp.get("tone", "navy"))
        _, h = sl.title(self.resolve(sp["headline"]), "headline", MX, 1.0, 9.6)
        raw = sp["items"]
        items = [{"title": t} for t in facts.site()["who_is_in_the_room"]] if raw == "@audience" else raw
        items = [{"title": i} if isinstance(i, str) else i for i in items]
        per_row = 4 if len(items) in (4, 7, 8) else 3 if len(items) in (3, 5, 6, 9) else len(items)
        gap = 0.4
        cw = (CW - gap * (per_row - 1)) / per_row
        y = max(3.1, 1.0 + h + 0.7)
        for i, it in enumerate(items):
            r, c = divmod(i, per_row)
            x, yy = MX + c * (cw + gap), y + r * 1.9
            sl.hline(x, yy, x + cw, color=sl.rule(True), name="Rule")
            sl.text(x, yy + 0.2, cw, 0.25, f"{i + 1:02d}", "label", sl.accent, name=f"Index {i + 1}")
            th = text_height_in(it["title"], "subtitle", cw)
            sl.text(x, yy + 0.6, cw, th + 0.05, it["title"], "subtitle", sl.ink, name=f"Entry {i + 1}")
            if it.get("text"):
                tt = self.resolve(it["text"])
                sl.text(x, yy + 0.7 + th, cw, text_height_in(tt, "datum", cw) + 0.05, tt, "datum", sl.soft, name=f"Entry {i + 1} text")
        return self._finish(sl, sp)

    # --- programme helpers ------------------------------------------------
    @staticmethod
    def _rows(day: int, track: str, frm=None, to=None, ids=None, breaks=True):
        d = facts.programme()[day - 1]
        rows = d["main"] if track == "main" else d["workshops"]
        if ids:
            valid = {r["id"] for r in rows}
            bad = [i for i in ids if i not in valid]
            if bad:
                raise ValueError(f"unknown session id(s) {bad} for day {day} {track}; valid: {sorted(valid)}")
            rows = [r for r in rows if r["id"] in ids]
        if frm:
            rows = [r for r in rows if r["time"] >= frm]
        if to:
            rows = [r for r in rows if r["time"] <= to]
        if not breaks:
            rows = [r for r in rows if r["kind"] != "break"]
        if not rows:
            raise ValueError(f"no programme rows for day {day} {track} {frm}-{to} {ids}")
        return d, rows

    # How a timed line is drawn is set per slide (see _slide_programme): speaker lines on/off, end times on/off.
    _speakers, _ends = True, False

    def _tcol(self):
        return 1.95 if self._ends else 1.45  # width of the gutter + time column

    def _time_label(self, r):
        return f"{r['time']}–{r['end']}" if self._ends and r.get("end") else r["time"]

    def _row_h(self, r, tw):
        h = text_height_in(r["title"], "session" if r["kind"] != "break" else "datum", tw)
        if r["subtitle"]:
            h += text_height_in(r["subtitle"], "datum", tw) + 0.04
        spk = facts.speaker_line(r) if self._speakers and r["kind"] != "break" else ""
        if spk:
            h += text_height_in(spk, "datum", tw) + 0.04
        return max(h, 0.3)

    def _draw_rows(self, sl: Slide, rows, x, y, w, limit=BOTTOM):
        """The timed line: a gold line down the gutter, a dot per entry, tabular time then title (then the
        speakers, as the agenda writes them). Returns (next_y, rows_left)."""
        tcol = self._tcol()
        time_x, title_x = x + 0.42, x + tcol
        tw = w - tcol
        time_w = tcol - 0.5
        top_y, drawn, last_dot = y, 0, None
        for idx, r in enumerate(rows):
            rh = self._row_h(r, tw)
            if y + rh > limit and drawn:
                break
            is_break = r["kind"] == "break"
            sl.oval(x + 0.01, y + 0.07, 0.12, PAL["gold"] if not is_break else mix(sl.ground, sl.soft, 0.45), name="Station dot")
            sl.text(time_x, y, time_w, 0.3, self._time_label(r), "datum", sl.accent if not is_break else sl.soft, name=f"Time {r['time']}")
            if is_break:
                sl.text(title_x, y, tw, rh + 0.03, r["title"], "datum", sl.soft, name=f"Break {r['time']}")
            else:
                paras = [[(r["title"], {})]]
                if r["subtitle"]:
                    paras.append([(r["subtitle"], {"role": "datum", "color": sl.soft})])
                spk = facts.speaker_line(r) if self._speakers else ""
                if spk:
                    paras.append([(spk, {"role": "datum", "color": sl.soft})])
                sl.text(title_x, y, tw, rh + 0.03, paras, "session", sl.ink, after=2, name=f"Session {r['time']}")
            last_dot = y + 0.13
            y += rh + 0.2
            drawn += 1
        if last_dot is not None and drawn:
            sl.vline(x + 0.07, top_y + 0.13, last_dot, color=PAL["gold"], pt=1.5, name="Programme line")
        return y, rows[drawn:]

    def _paginate(self, rows, tw, y0, limit=BOTTOM, gap=0.2, columns=1):
        """Split rows into the fewest pages that fit, then spread them evenly so no page is left nearly empty."""
        hs = [self._row_h(r, tw) + gap for r in rows]
        room = (limit - y0 + gap) * columns
        k, used = 1, 0.0
        for h in hs:
            if used + h > room and used:
                k, used = k + 1, 0.0
            used += h
        if k == 1:
            return [rows]
        target, pages, cur, used = sum(hs) / k, [], [], 0.0
        for r, h in zip(rows, hs):
            if cur and used + h / 2 > target and len(pages) < k - 1:
                pages.append(cur)
                cur, used = [], 0.0
            cur.append(r)
            used += h
        pages.append(cur)
        return pages

    def _split_columns(self, rows, tw, columns):
        """Cut a page's rows into `columns` runs of near-equal height, order preserved."""
        if columns == 1 or len(rows) < 2:
            return [rows]
        hs = [self._row_h(r, tw) for r in rows]
        half, acc, cut = sum(hs) / 2, 0.0, 1
        for i, h in enumerate(hs):
            acc += h
            if acc >= half:
                cut = i + 1
                break
        return [rows[:cut], rows[cut:]]

    def _slide_programme(self, sp):
        day, track = int(sp["day"]), sp.get("track", "main")
        columns = int(sp.get("columns", 1))
        if columns not in (1, 2):
            raise ValueError("columns must be 1 or 2")
        self._speakers, self._ends = bool(sp.get("speakers", True)), bool(sp.get("end_times", False))
        d, rows = self._rows(day, track, sp.get("from"), sp.get("to"), sp.get("ids"), sp.get("breaks", True))
        self._note_programme_flags(rows, track)
        marker = f"{d['day']} · {d['date']}" + (" · Workshops · Breakout room" if track == "workshops" else " · Main stage")
        colgap = 0.5
        colw = CW if columns == 1 else (CW - colgap) / 2
        pages = self._paginate(rows, colw - self._tcol(), 2.1, columns=columns)
        for page, chunk in enumerate(pages):
            sl = self.new_slide(sp.get("tone", "deep"))
            sl.text(MX, 0.72, 9, 0.25, marker + (f" · {page + 1} of {len(pages)}" if len(pages) > 1 else ""), "label", sl.accent, name="Day marker")
            sl.title(d["theme"], "title", MX, 1.12, 10.5)
            for ci, part in enumerate(self._split_columns(chunk, colw - self._tcol(), columns)):
                self._draw_rows(sl, part, MX + ci * (colw + colgap), 2.1, colw, limit=99)
            note = sp.get("note", facts.programme_disclaimer())  # the agenda page's own "Programme subject to change."
            if note and page == len(pages) - 1:
                sl.text(MX, 6.55, CW, 0.3, note, "datum", sl.soft, name="Note")
            self._finish(sl, sp if not page else {"notes": sp.get("notes")})

    def _note_programme_flags(self, rows, track):
        """Collect what the hand-off must carry whenever the agenda's rows are drawn; build() reports it once."""
        names = set(facts.agenda_only_names())
        acc = self.__dict__.setdefault("_agenda_flags", {"names": set(), "tba": set(), "overlap": set()})
        for r in rows:
            for s in r.get("speakers", []):
                nm = s["name"] if isinstance(s, dict) else s
                if (isinstance(s, dict) and s.get("org") == facts.TBA) or s == facts.TBA:
                    acc["tba"].add(f"{r['time']} {r['title']}")
                elif nm in names:
                    acc["names"].add(nm)
        ids = {r["id"] for r in rows}
        for day, a, b in facts.overlaps():
            if a["id"] in ids and b["id"] in ids:
                acc["overlap"].add(f"{day} {a['time']}–{a['end']} {a['title']!r} runs into {b['time']} {b['title']!r}")

    def _report_programme_flags(self):
        acc = self.__dict__.get("_agenda_flags")
        if not acc:
            return
        if acc["names"]:
            self.warn(0, "the programme shows people the agenda lists but the site's 26-speaker roster does not "
                         f"({', '.join(sorted(acc['names']))}) — confirm names and spellings before this leaves GCS")
        if acc["tba"]:
            self.warn(0, f"{len(acc['tba'])} slot(s) are 'to be announced' on the agenda and are drawn as such: " + "; ".join(sorted(acc["tba"])))
        for o in sorted(acc["overlap"]):
            self.warn(0, f"{o} — the site's THEMES.md says the organiser must confirm these times")

    # --- days (the two themes side by side, a few picked sessions each) ----
    def _slide_days(self, sp):
        sl = self.new_slide(sp.get("tone", "deep"))
        sl.title(self.resolve(sp.get("headline", "@site.programme_heading")), "headline", MX, 0.9, 10.2)
        gap = 0.6
        cw = (CW - gap) / 2
        pick = sp.get("pick", {})
        self._speakers, self._ends = bool(sp.get("speakers", False)), False
        for i, d in enumerate(facts.programme()):
            x = MX + i * (cw + gap)
            sl.hline(x, 2.55, x + cw, color=sl.rule(True), name="Rule")
            sl.text(x, 2.75, cw, 0.25, f"{d['day']} · {d['date']}", "label", sl.accent, name=f"Day {i + 1} marker")
            th = text_height_in(d["theme"], "title", cw)
            sl.text(x, 3.1, cw, th + 0.05, d["theme"], "title", sl.ink, name=f"Day {i + 1} theme")
            ids = pick.get(str(i + 1))
            if ids:
                _, rows = self._rows(i + 1, "main", ids=ids)
                self._draw_rows(sl, rows, x, 3.1 + th + 0.4, cw)
        return self._finish(sl, sp)

    # --- speakers --------------------------------------------------------
    def _slide_speakers(self, sp):
        names = sp.get("names", "all")
        roster = facts.speakers() if names == "all" else [facts.speaker(n) for n in names]
        per = int(sp.get("per_slide", 4 if len(roster) <= 4 else 8))
        if per not in (4, 8):
            raise ValueError("per_slide must be 4 or 8")
        pages = -(-len(roster) // per)
        base, extra = divmod(len(roster), pages)  # 26 at 8-up → 7, 7, 6, 6, not 8, 8, 8, 2
        chunks, i = [], 0
        for p in range(pages):
            size = base + (1 if p < extra else 0)
            chunks.append(roster[i:i + size])
            i += size
        for ci, chunk in enumerate(chunks):
            sl = self.new_slide(sp.get("tone", "deep"))
            head = self.resolve(sp.get("headline", "Speakers"))
            _, h = sl.title(head + (f" ({ci + 1}/{len(chunks)})" if len(chunks) > 1 else ""), "title" if per == 8 else "headline",
                            MX, 0.75, 10)
            if per == 4:
                gap = 0.3
                cw = (CW - gap * 3) / 4
                for i, s in enumerate(chunk):
                    x = MX + i * (cw + gap)
                    self._portrait(sl, s, x, 2.0, cw)
                    nh = text_height_in(s["display"], "subtitle", cw)
                    sl.text(x, 2.0 + cw + 0.2, cw, nh + 0.05, s["display"], "subtitle", sl.ink, name=f"Speaker {s['display']}")
                    sl.text(x, 2.0 + cw + 0.3 + nh, cw, 0.9, s["role"], "datum", sl.soft, name=f"Role {s['display']}")
            else:
                gap = 0.5
                cw = (CW - gap) / 2
                size, pitch = 1.1, 1.22
                for i, s in enumerate(chunk):
                    c, r = i % 2, i // 2
                    x, y = MX + c * (cw + gap), 1.6 + r * pitch
                    self._portrait(sl, s, x, y, size)
                    tx, tw = x + size + 0.25, cw - size - 0.25
                    nh = text_height_in(s["display"], "session", tw)
                    sl.text(tx, y, tw, nh + 0.03, s["display"], "session", sl.ink, name=f"Speaker {s['display']}")
                    rh = text_height_in(s["role"], "datum", tw, 1.1)
                    sl.text(tx, y + nh + 0.04, tw, rh + 0.03, s["role"], "datum", sl.soft, ls=1.1, name=f"Role {s['display']}")
                    if nh + 0.04 + rh > pitch - 0.02:
                        self.warn(sl.n, f"{s['display']}: name + role are {nh + 0.04 + rh:.2f}in tall, the row pitch is {pitch}in — use per_slide 4")
            self._finish(sl, sp if ci == 0 else {"notes": sp.get("notes")})

    def _portrait(self, sl: Slide, s, x, y, size):
        p = ASSETS / "speakers" / f"{s['slug']}.webp"
        if p.is_file():
            sl.picture(p, x, y, size, size, alt=f"Portrait of {s['display']}", focus=(0.5, 0.3), name=f"Portrait {s['display']}")
        else:  # initials, never a generated face
            sl.rect(x, y, size, size, PAL["dawn-low"], name=f"Portrait {s['display']}")
            sl.text(x, y, size, size, initials(s["display"]), "title", sl.accent, align="c", anchor="m", name="Initials")

    def _slide_speaker(self, sp):
        s = facts.speaker(sp["name"])
        sl = self.new_slide(sp.get("tone", "deep"))
        size = 4.0
        self._portrait(sl, s, MX, 1.4, size)
        x, w = MX + size + 0.7, W - MX - (MX + size + 0.7)
        _, h = sl.title(s["display"], "headline", x, 1.45, w)
        y = 1.45 + h + 0.15
        rh = text_height_in(s["role"], "lead", w)
        sl.text(x, y, w, rh + 0.05, s["role"], "lead", sl.ink, name="Role")
        y += rh + 0.35
        bh = text_height_in(s["bio"], "body", w)
        sl.text(x, y, w, bh + 0.05, s["bio"], "body", sl.soft, name="Bio")
        if y + bh > BOTTOM:
            self.warn(sl.n, f"{s['display']}: bio ends at {y + bh:.2f}in — too long for one slide; do not trim a verbatim bio, use a smaller portrait or split it")
        return self._finish(sl, sp)

    # --- rates -----------------------------------------------------------
    def _slide_rates(self, sp):
        sl = self.new_slide(sp.get("tone", "dawn"))
        r = facts.rates()
        _, h = sl.title(self.resolve(sp.get("headline", "Secure your place")), "headline", MX, 0.9, 9)
        inforce = facts.rate_on(self.as_of)
        if inforce is None:
            self.warn(sl.n, f"no published rate tier covers {self.as_of} — the slide shows all tiers with none marked current")
        y = 2.3
        colw = 7.0
        for t in r["tiers"]:
            current = inforce is not None and t["name"] == inforce["name"]
            sl.hline(MX, y, MX + colw, color=sl.rule(True), name="Rule")
            sl.text(MX, y + 0.2, 3.0, 0.4, t["name"], "subtitle", sl.ink if current else sl.soft, name=f"Tier {t['name']}")
            sl.text(MX, y + 0.62, 3.6, 0.3, t["when"], "datum", sl.soft, name=f"When {t['name']}")
            sl.text(MX + 3.9, y + 0.14, 3.1, 0.55, t["price"], "title", sl.accent if current else sl.soft, align="r", name=f"Price {t['name']}")
            if current:
                sl.text(MX + 3.9, y + 0.7, 3.1, 0.25, "Current rate", "label", sl.accent, align="r", name="Current rate")
            y += 1.18
        x2, w2 = 8.6, W - MX - 8.6
        sl.text(x2, 2.3, w2, 0.25, "Included", "label", sl.accent, name="Included marker")
        incl = r["included"]
        ih = sum(text_height_in(i, "body", w2) for i in incl) + 0.16 * (len(incl) - 1)
        sl.text(x2, 2.75, w2, ih + 0.05, incl, "body", sl.ink, after=11, name="Included")
        sl.text(x2, 2.75 + ih + 0.3, w2, text_height_in(r["limit"], "datum", w2) + 0.05, r["limit"], "datum", sl.soft, name="Limit")
        self._button(sl, MX, 6.0, r["cta"], sp.get("url"))
        return self._finish(sl, sp)

    def _button(self, sl: Slide, x, y, label, url=None):
        """The one lit block: gold ground, navy text, square. Linked only when the URL is confirmed."""
        bw = text_width_pt(label.upper(), "body", 12, 1.92) / 72 + 0.8
        # on a night ground the call to action is the one lit gold block; on paper it is navy with white text
        b = sl.rect(x, y, bw, 0.56, PAL["navy"] if sl.paper else PAL["gold"], name="Call to action")
        sl.text(x, y, bw, 0.56, label, "label", PAL["white"] if sl.paper else PAL["navy"], align="c", anchor="m", name="Call to action label",
                link=url if url and "url" in self.confirmed else None)
        if url and "url" not in self.confirmed:
            self.warn(sl.n, "a URL was given but meta.confirmed does not include 'url' — the button is not linked. Confirm the URL (the site sat behind a staging gate) and re-run")
        return b

    # --- section ---------------------------------------------------------
    def _slide_section(self, sp):
        sl = self.new_slide(sp.get("tone", "deep"))
        if sp.get("number"):
            sl.text(MX, 2.6, 4, 0.25, sp["number"], "label", sl.accent, name="Section number")
        sl.title(self.resolve(sp["headline"]), "headline", MX, 3.0, 9.4)
        return self._finish(sl, sp)

    # --- closing ---------------------------------------------------------
    def _slide_closing(self, sp):
        """The back cover: dark blue, the cover's masthead and marks, one call to action."""
        sl = self.new_slide(sp.get("tone", "navy"), cover=True)
        sl.logo("wm-horizontal", MX, 0.62, w=3.5, alt="Wealth & Migration Summit")
        service = self.resolve(sp.get("service_line", "@site.service_line"))
        sl.rect(MX, 1.775, 0.1, 0.1, PAL["gold"], name="Lamp")
        sl.text(MX + 0.28, 1.7, 7, 0.25, service, "label", sl.accent, name="Service line")
        _, h = sl.title(self.resolve(sp["headline"]), "display", MX, 2.25, 9.6)
        y = 2.25 + h + 0.35
        if sp.get("lead"):
            lead = self.resolve(sp["lead"])
            lh = text_height_in(lead, "lead", 7.4)
            sl.text(MX, y, 7.4, lh + 0.05, lead, "lead", sl.soft, name="Lead")
            y += lh + 0.4
        by = max(y, 4.4)
        if sp.get("cta", "@rates.cta") is not False:  # "cta": false = no button, e.g. for delegates who already hold a place
            self._button(sl, MX, by, self.resolve(sp.get("cta", "@rates.cta")), sp.get("url"))
            by += 0.85
        if sp.get("contact"):
            sl.text(MX, by, 8, 0.3, sp["contact"], "datum", sl.soft, name="Contact")
            by += 0.4
        if by > 6.3:
            self.warn(sl.n, f"back cover stack ends at {by:.2f}in, into the logo zone — shorten the headline or lead")
        self._marks(sl, sp.get("marks", ["gcs", "cnn"]))
        return self._finish(sl, sp)


SLIDE_TYPES = sorted(n[7:] for n in dir(Deck) if n.startswith("_slide_"))


# ----------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("spec", help="deck spec (JSON)")
    ap.add_argument("out", help="output .pptx")
    ap.add_argument("--fonts", choices=sorted(FONT_SETS), help="office (default) = Cambria/Arial; brand = Caladea/Archivo, only where installed")
    ap.add_argument("--as-of", help="date the deck is valid for (default: today) — decides which rate tier is current")
    ap.add_argument("--no-check", action="store_true", help="skip the fact/layout lint that normally runs after the build")
    args = ap.parse_args(argv)

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    meta = spec.get("meta", {})
    deck = Deck(fonts=args.fonts or meta.get("fonts", "office"), as_of=args.as_of or meta.get("as_of"),
                lang=meta.get("lang", "en-GB"), title=meta.get("title", "Wealth & Migration Summit 2026"),
                author=meta.get("author", "Global Citizen Solutions"), confirmed=meta.get("confirmed"))
    deck.build(spec)
    out = deck.save(args.out)
    print(f"built {out}  ({len(deck.prs.slides)} slides, fonts: {deck.fonts}, as of {deck.as_of})")
    for n, msg in deck.warnings:
        print(f"  WARN  {'slide ' + str(n) if n else 'deck'}: {msg}")
    if args.no_check:
        return 0
    import check_deck
    return check_deck.run(out, as_of=deck.as_of, confirmed=deck.confirmed, spec=spec)


if __name__ == "__main__":
    sys.exit(main())
